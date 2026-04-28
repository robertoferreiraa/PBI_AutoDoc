"""
documenta.py — Geração de documentação Word e Excel + orquestração das chamadas LLM
"""

import io
import json
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from llm_client import call_llm, call_llm_text, chunk_text_by_tokens
from prompts import (
    SYSTEM_PROMPT,
    prompt_chat,
    prompt_documentacao_completa,
    prompt_fontes,
    prompt_medidas,
)

# ---------------------------------------------------------------------------
# Constantes de estilo
# ---------------------------------------------------------------------------

HEADER_COLOR = "1A3A6B"     # Azul escuro
HEADER_TEXT_COLOR = RGBColor(255, 255, 255)
ACCENT_COLOR = RGBColor(26, 58, 107)
MAX_TOKENS_CHUNK = 3000
MAX_TOKENS_OUTPUT = 8192


# ---------------------------------------------------------------------------
# Orquestração da documentação via LLM
# ---------------------------------------------------------------------------

def gerar_documentacao(
    data: dict,
    modelo: str,
    language: str = "Português",
    progress_callback=None,
) -> dict:
    """
    Orquestra todas as chamadas LLM para gerar a documentação completa.

    Parâmetros
    ----------
    data : dict  — resultado do extractor.extract_from_file()
    modelo : str — nome do modelo LiteLLM
    language : str — idioma das descrições
    progress_callback : callable(step: int, total: int, msg: str) | None

    Retorna
    -------
    dict com:
        - info: dict  (dados gerais do relatório)
        - tables: list[dict]
        - measures: list[dict]
        - sources: list[dict]
    """
    report_name = data["report_name"]
    tables_df: pd.DataFrame = data["tables_df"]
    measures_df: pd.DataFrame = data["measures_df"]

    def _progress(step, total, msg):
        if progress_callback:
            progress_callback(step, total, msg)

    # ---- 1. Documentação geral (info + tabelas + fontes numa só chamada) ----
    _progress(1, 4, "Gerando documentação geral do relatório...")

    tabelas_texto = tables_df["NomeTabela"].tolist() if not tables_df.empty else []
    fontes_texto = []
    if not tables_df.empty and "FonteDados" in tables_df.columns:
        for _, row in tables_df.iterrows():
            if row.get("FonteDados"):
                fontes_texto.append(
                    f"Tabela: {row['NomeTabela']}\nFonte M:\n{row['FonteDados'][:500]}"
                )

    user_prompt_geral = (
        prompt_documentacao_completa(language)
        + f"\n\nNome do relatório: {report_name}\n"
        + f"\nTabelas:\n" + "\n".join(f"- {t}" for t in tabelas_texto)
        + f"\n\nFontes de dados (primeiras):\n" + "\n---\n".join(fontes_texto[:5])
    )

    resultado_geral = call_llm(
        modelo, SYSTEM_PROMPT, user_prompt_geral, max_tokens=MAX_TOKENS_OUTPUT
    )

    info = resultado_geral.get("Relatorio", {
        "Titulo": report_name,
        "Descricao": "Documentação gerada automaticamente.",
        "Principais_KPIs_e_Metricas": [],
        "Publico_Alvo": "Analistas de dados",
        "Exemplos_de_Uso": [],
    })
    tables_doc = resultado_geral.get("Tabelas_do_Relatorio", [])
    sources_doc = resultado_geral.get("Fontes_de_Dados", [])

    # ---- 2. Medidas em chunks ----
    _progress(2, 4, "Documentando medidas DAX...")

    measures_doc: list[dict] = []
    if not measures_df.empty:
        segments = [
            f"Nome: {row['NomeMedida']}\nExpressão DAX: {row['ExpressaoMedida']}"
            for _, row in measures_df.iterrows()
        ]
        chunks = chunk_text_by_tokens(segments, max_tokens=MAX_TOKENS_CHUNK)

        for i, chunk in enumerate(chunks):
            _progress(2, 4, f"Documentando medidas — parte {i+1}/{len(chunks)}...")
            user_prompt_med = prompt_medidas(language) + f"\n\n{chunk}"
            result = call_llm(modelo, SYSTEM_PROMPT, user_prompt_med, max_tokens=MAX_TOKENS_OUTPUT)
            measures_doc.extend(result.get("Medidas_do_Relatorio", []))

    # ---- 3. Fontes de dados em chunks (detalhado) ----
    _progress(3, 4, "Documentando fontes de dados...")

    if not tables_df.empty and "FonteDados" in tables_df.columns:
        fontes_segments = [
            f"Tabela: {row['NomeTabela']}\nFonte M:\n{row['FonteDados']}"
            for _, row in tables_df.iterrows()
            if row.get("FonteDados")
        ]
        if fontes_segments:
            chunks_f = chunk_text_by_tokens(fontes_segments, max_tokens=MAX_TOKENS_CHUNK)
            sources_doc_detalhado: list[dict] = []
            for i, chunk in enumerate(chunks_f):
                _progress(3, 4, f"Documentando fontes — parte {i+1}/{len(chunks_f)}...")
                user_prompt_src = prompt_fontes(language) + f"\n\n{chunk}"
                result = call_llm(modelo, SYSTEM_PROMPT, user_prompt_src, max_tokens=MAX_TOKENS_OUTPUT)
                sources_doc_detalhado.extend(result.get("Fontes_de_Dados", []))
            if sources_doc_detalhado:
                sources_doc = sources_doc_detalhado

    _progress(4, 4, "Documentação concluída!")

    return {
        "info": info,
        "tables": tables_doc,
        "measures": measures_doc,
        "sources": sources_doc,
    }


# ---------------------------------------------------------------------------
# Chat interativo
# ---------------------------------------------------------------------------

def chat_sobre_relatorio(
    pergunta: str,
    data: dict,
    doc_resultado: dict,
    modelo: str,
    language: str = "Português",
) -> str:
    """Responde perguntas sobre o relatório usando o contexto dos metadados."""
    tables_df: pd.DataFrame = data["tables_df"]
    measures_df: pd.DataFrame = data["measures_df"]

    # Monta contexto resumido
    tabelas = tables_df["NomeTabela"].tolist() if not tables_df.empty else []
    medidas = measures_df["NomeMedida"].tolist() if not measures_df.empty else []

    context = f"""Relatório: {data['report_name']}

Tabelas ({len(tabelas)}): {', '.join(tabelas[:30])}

Medidas ({len(medidas)}): {', '.join(medidas[:30])}

Documentação gerada:
{json.dumps(doc_resultado, ensure_ascii=False, indent=2)[:3000]}
"""
    system = prompt_chat(context, language)
    return call_llm_text(modelo, system, pergunta, max_tokens=2048, temperature=0.3)


# ---------------------------------------------------------------------------
# Utilitários Word
# ---------------------------------------------------------------------------

def _add_heading(doc: Document, text: str, level: int = 1):
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.color.rgb = ACCENT_COLOR
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT


def _add_bullet(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(str(item))
        run.font.size = Pt(11)


def _style_header_cell(cell):
    """Aplica cor de fundo e texto branco ao cabeçalho da tabela."""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), HEADER_COLOR)
    tc_pr.append(shd)
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = HEADER_TEXT_COLOR
            run.bold = True


def _add_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            tc = cell._element
            tc_pr = tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{side}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:space"), "0")
                border.set(qn("w:color"), "CCCCCC")
                borders.append(border)
            tc_pr.append(borders)


def _set_col_widths(table, widths: list[float]):
    for i, col in enumerate(table.columns):
        if i < len(widths):
            for cell in col.cells:
                cell.width = Inches(widths[i])


# ---------------------------------------------------------------------------
# Geração do Word
# ---------------------------------------------------------------------------

def generate_docx(
    doc_resultado: dict,
    data: dict,
    modelo: str,
    language: str = "Português",
) -> io.BytesIO:
    """Gera o documento Word e retorna como BytesIO."""
    doc = Document()
    info = doc_resultado["info"]
    measures_df: pd.DataFrame = data["measures_df"]
    df_rels = data.get("df_relationships")
    columns_df: pd.DataFrame = data["columns_df"]

    # --- Cabeçalho ---
    title_para = doc.add_paragraph()
    title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title_para.add_run("📊 Documentação Power BI")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = ACCENT_COLOR

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle.add_run(f"Gerado por: {modelo}  |  {datetime.now().strftime('%d/%m/%Y %H:%M')}").italic = True

    doc.add_paragraph()

    # --- Informações do Relatório ---
    _add_heading(doc, "Informações do Relatório", 1)
    doc.add_paragraph(f"Título: {info.get('Titulo', data['report_name'])}", style="Body Text")
    doc.add_paragraph()

    _add_heading(doc, "Descrição", 2)
    doc.add_paragraph(info.get("Descricao", ""), style="Body Text")

    _add_heading(doc, "Principais KPIs e Métricas", 2)
    _add_bullet(doc, info.get("Principais_KPIs_e_Metricas", []))

    _add_heading(doc, "Público-Alvo", 2)
    doc.add_paragraph(info.get("Publico_Alvo", ""), style="Body Text")

    _add_heading(doc, "Exemplos de Uso", 2)
    _add_bullet(doc, info.get("Exemplos_de_Uso", []))

    doc.add_paragraph()

    # --- Tabelas ---
    _add_heading(doc, "Tabelas do Relatório", 1)
    tables = doc_resultado.get("tables", [])
    if tables:
        tbl = doc.add_table(rows=1, cols=2)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Tabela"
        hdr[1].text = "Descrição"
        for cell in hdr:
            _style_header_cell(cell)
        _set_col_widths(tbl, [2.5, 6.0])
        for t in tables:
            row = tbl.add_row().cells
            row[0].text = str(t.get("Nome", ""))
            row[1].text = str(t.get("Descricao", ""))
        _add_table_borders(tbl)
    else:
        doc.add_paragraph("Nenhuma tabela documentada.", style="Body Text")

    doc.add_paragraph()

    # --- Medidas ---
    _add_heading(doc, "Medidas DAX", 1)
    measures = doc_resultado.get("measures", [])
    if measures:
        tbl = doc.add_table(rows=1, cols=3)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Medida"
        hdr[1].text = "Descrição"
        hdr[2].text = "Fórmula DAX"
        for cell in hdr:
            _style_header_cell(cell)
        _set_col_widths(tbl, [2.0, 4.5, 3.0])
        for m in measures:
            nome = m.get("Nome", "")
            descr = m.get("Descricao", "")
            # Busca expressão DAX no dataframe original
            dax = ""
            if not measures_df.empty:
                match = measures_df.loc[measures_df["NomeMedida"] == nome, "ExpressaoMedida"]
                dax = match.values[0] if not match.empty else ""
            row = tbl.add_row().cells
            row[0].text = nome
            row[1].text = descr
            row[2].text = dax
        _add_table_borders(tbl)
    else:
        doc.add_paragraph("Nenhuma medida documentada.", style="Body Text")

    doc.add_paragraph()

    # --- Fontes de Dados ---
    _add_heading(doc, "Fontes de Dados", 1)
    sources = doc_resultado.get("sources", [])
    if sources:
        tbl = doc.add_table(rows=1, cols=3)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Fonte"
        hdr[1].text = "Descrição"
        hdr[2].text = "Tabelas"
        for cell in hdr:
            _style_header_cell(cell)
        _set_col_widths(tbl, [2.5, 5.0, 2.0])
        for s in sources:
            tabelas_m = s.get("Tabelas_Contidas_no_M", [])
            if isinstance(tabelas_m, list):
                tabelas_str = ", ".join(tabelas_m)
            else:
                tabelas_str = str(tabelas_m)
            row = tbl.add_row().cells
            row[0].text = str(s.get("Nome", ""))
            row[1].text = str(s.get("Descricao", ""))
            row[2].text = tabelas_str
        _add_table_borders(tbl)
    else:
        doc.add_paragraph("Nenhuma fonte de dados documentada.", style="Body Text")

    doc.add_paragraph()

    # --- Colunas ---
    _add_heading(doc, "Colunas das Tabelas", 1)
    if not columns_df.empty:
        tbl = doc.add_table(rows=1, cols=4)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Tabela"
        hdr[1].text = "Coluna"
        hdr[2].text = "Tipo"
        hdr[3].text = "Expressão"
        for cell in hdr:
            _style_header_cell(cell)
        _set_col_widths(tbl, [2.0, 2.5, 1.5, 3.5])
        for _, row in columns_df.iterrows():
            r = tbl.add_row().cells
            r[0].text = str(row.get("NomeTabela", ""))
            r[1].text = str(row.get("NomeColuna", ""))
            r[2].text = str(row.get("TipoDadoColuna", ""))
            r[3].text = str(row.get("ExpressaoColuna", "") or "")
        _add_table_borders(tbl)
    else:
        doc.add_paragraph("Nenhuma coluna encontrada.", style="Body Text")

    doc.add_paragraph()

    # --- Relacionamentos ---
    if df_rels is not None and not df_rels.empty:
        _add_heading(doc, "Relacionamentos", 1)
        tbl = doc.add_table(rows=1, cols=4)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Tabela Origem"
        hdr[1].text = "Coluna Origem"
        hdr[2].text = "Tabela Destino"
        hdr[3].text = "Coluna Destino"
        for cell in hdr:
            _style_header_cell(cell)
        _set_col_widths(tbl, [2.5, 2.5, 2.5, 2.5])
        for _, row in df_rels.iterrows():
            r = tbl.add_row().cells
            r[0].text = str(row.get("FromTable", ""))
            r[1].text = str(row.get("FromColumn", ""))
            r[2].text = str(row.get("ToTable", ""))
            r[3].text = str(row.get("ToColumn", ""))
        _add_table_borders(tbl)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Geração do Excel
# ---------------------------------------------------------------------------

def generate_excel(doc_resultado: dict, data: dict) -> io.BytesIO:
    """Gera o arquivo Excel e retorna como BytesIO."""
    buf = io.BytesIO()
    measures_df: pd.DataFrame = data["measures_df"]
    df_rels = data.get("df_relationships")
    columns_df: pd.DataFrame = data["columns_df"]
    tables_df: pd.DataFrame = data["tables_df"]

    info = doc_resultado["info"]
    tables = doc_resultado.get("tables", [])
    measures = doc_resultado.get("measures", [])
    sources = doc_resultado.get("sources", [])

    with pd.ExcelWriter(buf, engine="xlsxwriter") as writer:
        wb = writer.book

        # Formatos
        fmt_header = wb.add_format({
            "bold": True, "bg_color": f"#{HEADER_COLOR}", "font_color": "white",
            "border": 1, "text_wrap": True, "valign": "top"
        })
        fmt_cell = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
        fmt_title = wb.add_format({"bold": True, "font_size": 14, "font_color": f"#{HEADER_COLOR}"})

        def write_sheet(name: str, df: pd.DataFrame):
            if df.empty:
                return
            df.to_excel(writer, sheet_name=name, index=False, startrow=1)
            ws = writer.sheets[name]
            ws.write(0, 0, name, fmt_title)
            for col_i, col_name in enumerate(df.columns):
                ws.write(1, col_i, col_name, fmt_header)
            for row_i, row in enumerate(df.values):
                for col_i, val in enumerate(row):
                    ws.write(row_i + 2, col_i, str(val) if val is not None else "", fmt_cell)
            ws.set_column(0, len(df.columns) - 1, 30)

        # Aba: Info
        info_df = pd.DataFrame([
            {"Campo": k, "Valor": str(v) if not isinstance(v, list) else ", ".join(v)}
            for k, v in info.items()
        ])
        write_sheet("Informações", info_df)

        # Aba: Tabelas
        if tables:
            write_sheet("Tabelas", pd.DataFrame(tables))

        # Aba: Medidas
        if measures:
            df_med = pd.DataFrame(measures)
            # Mescla com as expressões DAX originais
            if not measures_df.empty:
                df_med = df_med.merge(
                    measures_df[["NomeMedida", "ExpressaoMedida"]],
                    left_on="Nome",
                    right_on="NomeMedida",
                    how="left",
                )
                df_med = df_med.rename(columns={"ExpressaoMedida": "Fórmula DAX"})
                df_med = df_med.drop(columns=["NomeMedida"], errors="ignore")
            write_sheet("Medidas", df_med)

        # Aba: Fontes
        if sources:
            df_src = pd.DataFrame(sources)
            if "Tabelas_Contidas_no_M" in df_src.columns:
                df_src["Tabelas_Contidas_no_M"] = df_src["Tabelas_Contidas_no_M"].apply(
                    lambda x: ", ".join(x) if isinstance(x, list) else str(x)
                )
            write_sheet("Fontes de Dados", df_src)

        # Aba: Colunas
        if not columns_df.empty:
            write_sheet("Colunas", columns_df)

        # Aba: Relacionamentos
        if df_rels is not None and not df_rels.empty:
            write_sheet("Relacionamentos", df_rels)

    buf.seek(0)
    return buf
